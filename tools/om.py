#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
om.py — doc file order noi dung, chuan hoa, chay Cong A, suy ra cac o.

Dung:
    python tools/om.py <file.csv|file.xlsx> [--sheet TEN] [--out out/phieu.csv]
    python tools/om.py --gid <GID> --sheet-id <SPREADSHEET_ID> [--out ...]

Script nay KHONG viet kich ban. No lam B1-B3 (nap, chuan hoa, cong A) va
dien san cac o suy ra duoc, de agent chi phai lam phan sang tao: B4-B7.

Doi chieu: knowledge/om-B-doc-dau-vao.md
"""

from __future__ import annotations

import argparse
import csv
import io
import re
import sys
import unicodedata
import urllib.request
from pathlib import Path

# ---------------------------------------------------------------- tri thuc tra cuu

# om-B §2 — chuan hoa chan dung
CHAN_DUNG = {
    "khuyet diem": ("CD1", "Khuyet diem"),
    "hot face": ("CD2", "Hot Face"),
    "sua lai": ("CD3", "Sua lai"),
    "nam gioi": ("CD4", "Nam gioi"),
    "phong thuy": ("CD5", "Phong thuy"),
    "kho benh ly": ("CD6", "Kho, benh ly"),
    # om-J §1 — chuyen khoa MAT: sup mi la benh ly -> CD6; lao hoa -> CD5 (MA DANG LECH)
    "sup mi": ("CD6", "Sup mi & benh ly"),
    "lao hoa": ("CD5", "Lao hoa"),
}

# om-B §3 — nac 5A -> nhiet
NHIET = {"A1": "Cold", "A2": "Warm", "A3": "Warm", "A4": "Hot", "A5": "Warm"}

# om-B §3 — nac x camp dich -> o camp
O_CAMP = {
    ("A1", "VIEW"): "V1", ("A1", "PERF"): "CHAN",
    ("A2", "VIEW"): "V2", ("A2", "PERF"): "P2",
    ("A3", "VIEW"): "V3", ("A3", "PERF"): "P3",
    ("A4", "VIEW"): "V4", ("A4", "PERF"): "P4",
    ("A5", "VIEW"): "V5", ("A5", "PERF"): "P5",
}

# om-F §1 — format cho phep theo o
FORMAT_THEO_O = {
    "V1": [], "V2": ["L06", "L02", "L10"], "V3": ["L09", "L08", "L14"],
    "V4": ["L15", "L11"], "V5": ["L11", "L13"],
    "P2": ["L06", "L10", "L07"], "P3": ["L08", "L09", "L04", "L14"],
    "P4": ["L04", "L15", "L11"], "P5": ["L06", "L11"],
}

# om-F §2 — kho tep
TEP = {
    "V1": ("LAL-CHOT · LAL-SDT · LAL-inb · BROAD-GEO", "CUST"),
    "V2": ("VV25-DV · VV25-CUM", "MSG · BOOKED · CUST"),
    "V3": ("VV95-N1 · VV95-N2", "MSG · LEAD-NB · BOOKED · CUST"),
    "V4": ("BOOKED-NS", "CUST"),
    "V5": ("CUST", "-"),
    "P2": ("LAL-SDT · LAL-inb · BROAD-GEO", "ENG · MSG · BOOKED · CUST"),
    "P3": ("VV25-DV · VV25-CUM · VV95-N2 · ENG", "MSG · LEAD-NB · BOOKED · CUST"),
    "P4": ("LEAD-NB · MSG · WEB-BK", "BOOKED · CUST"),
    "P5": ("CUST", "BOOKED (dang dieu tri)"),
}

# om-F §3 — nguong
NGUONG = {
    "V1": ("<=150.000d / 1.000 ThruPlay", "1.000 ThruPlay"),
    "V2": ("<=170.000d / 1.000 ThruPlay", "1.000 ThruPlay"),
    "V3": ("<=170.000d / 1.000 ThruPlay", "1.000 ThruPlay"),
    "V4": ("<=170.000d / 1.000 ThruPlay", "1.000 ThruPlay"),
    "V5": ("<=150.000d / 1.000 ThruPlay", "1.000 ThruPlay"),
    "P2": ("<=400.000d / hoi thoai", "Hoi thoai"),
    "P3": ("<=1.150.000d / SDT hop le", "Lead SDT hop le"),
    "P4": ("<=2.650.000d / lich hen", "Lich hen"),
    "P5": ("[CHO XAC NHAN]", "Hoi thoai cross-sell"),
}

# om-D — thu vien format: ma -> (ten, thoi luong, boi canh, chi so do, phut quay, can KH that)
FORMAT = {
    "L01": ("Bac si - Talk hoi dap", "25-35s", "S-STUDIO", "Giu chan 3s · share", 9, False),
    "L02": ("Bac si - Talk chu de", "40-60s", "S-KHAM", "Xem het · binh luan chuyen mon", 9, False),
    "L03": ("Bac si - Phan tich", "30-45s", "S-KHAM", "Xem het · luu", 18, False),
    "L04": ("Khach hang - Truoc/sau", "10-15s", "S-STUDIO", "CTR · ty le inbox", 11, True),
    "L05": ("Khach hang - Bien doi", "15-25s", "S-STUDIO", "Binh luan / 1.000 view", 18, True),
    "L06": ("Khach hang - Review", "<=45s", "S-STUDIO", "Luu · share", 38, True),
    "L07": ("Khach hang - Hanh trinh", "45-90s", "S-SANH", "LUU (cao nhat he)", 38, True),
    "L08": ("Khach hang - Ket qua", "20-35s", "S-PHUCHOI", "Ty le inbox (cao nhat he)", 18, False),
    "L09": ("Dich vu - Cong nghe", "15-30s", "S-PHUCHOI", "Binh luan tich cuc · xem het", 18, False),
    "L10": ("Dich vu - Quy trinh", "25-40s", "S-TIENME", "Binh luan tich cuc · luu", 25, False),
    "L11": ("Dich vu - Gia/Sales/Chinh sach", "10-20s", "S-STUDIO", "Share · binh luan tich cuc", 11, False),
    "L12": ("Local - Co so/Bac si/Cong nghe", "20-35s", "S-STUDIO", "LUU · giam huy lich", 25, False),
}
FORMAT_CHUA_DINH_NGHIA = {"L13", "L14", "L15", "L16"}

# om-D §3 — dao cu bat buoc theo format
DAO_CU = {
    "L01": "Teleprompter · mic khong day · blouse co dinh",
    "L02": "Teleprompter · mic khong day · blouse co dinh",
    "L03": "Anh/CT DA CHE MAT · man hinh 27\" · but chi · bang viet kinh",
    "L04": "Anh/video before dung goc · thuoc",
    "L05": "Bo anh moc ngay 1-3-7-14-30 CUNG MOT khach · do hoa moc ngay",
    "L06": "BAN CONSENT DA KY · mic cai ao cho khach · gimbal",
    "L07": "BAN CONSENT DA KY · mic cai ao cho khach · gimbal",
    "L08": "Mau vat lieu con seal · thuoc · anh before dung goc",
    "L09": "Thiet bi/may moc that · do hoa so · mo hinh minh hoa",
    "L10": "Checklist in san · do hoa 5 buoc",
    "L11": "BANG GIA HIEN HANH ban in · do hoa so suat",
    "L12": "Canh quan co so · dong phuc ekip",
}

# om-D §1 — nhan vat can co mat theo format
NHAN_VAT = {
    "L01": "Bac si", "L02": "Bac si", "L03": "Bac si",
    "L04": "Khach hang", "L05": "Khach hang",
    "L06": "Khach hang", "L07": "Khach hang + dieu duong",
    "L08": "Bac si + khach hang", "L09": "Bac si + khach hang",
    "L10": "Dieu duong", "L11": "Bac si hoac MC", "L12": "Ekip",
}

# om-B §1 col 8 — ten format tu do trong order -> ma L
TEN_FORMAT = {
    "talk hoi dap": "L01", "talk chu de": "L02", "phan tich": "L03",
    "truoc sau": "L04", "truoc/ sau": "L04", "bien doi": "L05",
    "review": "L06", "hanh trinh": "L07", "ket qua": "L08",
    "cong nghe": "L09", "quy trinh": "L10",
    "gia": "L11", "sales": "L11", "chinh sach": "L11", "local": "L12",
}

# om-C §1 — noi so chi phoi + thong diep thang
CHAN_DUNG_TRI_THUC = {
    "CD1": ("So mui gia/do, khong hop mat; so bien chung phai lam lai; so bi 'lam ho'",
            "Mot lan nang mui - dep tu nhien tron doi - an toan la tren het"),
    "CD2": ("So mui khong len hinh dep; so sung lau khong livestream/chup job duoc",
            "Nang mui - len hinh auto xinh, 7 ngay hoi phuc"),
    "CD3": ("So hong lan nua; so khong ai chiu trach nhiem; so cang sua cang xau",
            "Sua mui hong - ca cang kho, bac si cang chuyen sau"),
    "CD4": ("So bi che 'dan ong di tham my'; so mui cao qua thanh gia, nu tinh",
            "Mui chuan men - dep tu nhien, khong ai biet da lam"),
    "CD5": ("So pha mat loc hien tai; so mui hech lo lo mui ('hao tai ton cua')",
            "Doi dang mui - mo cung tai: chuan nhan tuong"),
    "CD6": ("So phau thuat that bai (lon nhat); so bien chung anh huong chuc nang tho",
            "Tai sinh dang mui - chua lanh moi khuyet diem, tra lai tu nhien"),
}

# om-J §2 — chan dung MAT (phan lon la SUY RA, xem om-J §0)
CHAN_DUNG_MAT = {
    "CD1": ("SO CHI PHI PHAT SINH (noi so so 1); so seo goc mat trong; so sung lau nghi lam",
            "Cat mi dung tinh trang - nep mi ro net, hoi phuc nhanh, gia tron goi khong phat sinh"),
    "CD2": ("So mat len hinh thay nep nhan; so sung lo job; so lam xong mat nhan dien",
            "Mat khong gia di theo nghe"),
    "CD3": ("So sua nua van khong dep; so da tieu tien vo ich; so khong ai chi ra van de that",
            "Cat mi don thuan khong giai quyet tat ca - bac si chi ra van de that cua ban"),
    "CD4": ("SO BI DANH GIA VI DI THAM MY; so mat net nam tinh",
            "Mat co hon, tam nhin thoang hon - khong ai biet da lam"),
    "CD5": ("SO NGUY HIEM VI TUOI CAO, BENH NEN; so hoi phuc lau",
            "Trong tuoi hon, do met moi - khong phai tre nhu gai 20"),
    "CD6": ("So khong sua duoc; so anh huong chuc nang mat; so hai ben van khong deu",
            "Mat mo binh thuong, hai ben deu"),
}

# om-J §3 — CTA chuan cua nac, chuyen khoa MAT
CTA_NAC_MAT = {
    "A1": "Xem mat ban thuoc kieu nao",
    "A2": "Gui anh mat - bac si chi ra ban can chinh gi",
    "A3": "Inbox nhan bao gia theo tinh trang",
    "A4": "Inbox giu suat / giu lich",
    "A5": "Luu video - nhan neu can ho tro hau phau",
}

# om-C §3 — CTA chuan cua nac
CTA_NAC = {
    "A1": "Xem hanh trinh lot xac 7 ngay",
    "A2": "Gui anh - bac si tu van dang mui hop mat",
    "A3": "Dat lich tu van truc tiep voi bac si",
    "A4": "Inbox giu lich / giu suat",
    "A5": "Luu video - nhan neu can ho tro hau phau",
}

# om-C §5 — 7 kieu hook KEM
KIEU_KEM = [
    ("H1", "tinh tu chung chung", [
        "van nguoi me", "slay", "net cang", "thang hang nhan sac", "xuat sac hon",
        "dep nhu y", "chan ai", "dinh cao", "hoan my"]),
    ("H2", "mo bang gia", [r"^\s*\d+\s*tr", r"chi\s+\d+\s*tr", r"chi\s+tu\s+\d+"]),
    ("H5", "hua moc thoi gian / ket qua", [
        "dep ngay", "bao lau dep", "cam ket 100", "khoi hoan toan",
        "dep ngay sau", "lot xac ngay"]),
    ("H7", "tu giat gan", ["chan dong", "soc", "khong the tin", "bat ngo 100"]),
]
# H3 (mo ho), H4 (qua ngan) kiem bang do dai + co tu chi tinh trang
TU_TINH_TRANG = [
    "tep", "thap", "hech", "ngan", "go", "lech", "co rut", "bong do", "viem",
    "lo song", "bien dang", "di tat", "sup", "bong mo", "da du", "chung",
    "filler", "hong", "khuyet", "to", "canh mui", "vach ngan", "nghet",
]
TU_KY_THUAT = [
    "cau truc", "ban cau truc", "sun suon", "sun tai", "sun nhan tao", "can co",
    "tiem tan", "tai cau truc", "gia co", "tru mui", "sieu am", "ct", "noi soi",
]
TU_BANG_CHUNG = [
    "nam sau", "nam roi", "ca that", "ca nay", "boc seal", "truoc sau",
    "hanh trinh", "ngay 1", "theo doi", "tai kham", "1 ca", "mot ca",
]

# om-G §2 — co phap che
CO_DO = [
    ("tron doi", "L03 L04", "chinh sach bao hanh phai doi chieu van ban"),
    ("cam ket", "L03", "cam ket ket qua"),
    ("dep ngay", "L03", "cam ket ket qua, sai thuc te hau phau"),
    ("100%", "L03 L04", "khang dinh tuyet doi / con so khong nguon"),
    ("doi van", "L03 L04", "hua hieu qua phong thuy"),
    ("tai loc", "L03 L04", "hua hieu qua phong thuy"),
    ("giai han", "L03 L04", "hua hieu qua phong thuy"),
    ("bac si chi", "L01 L02", "ten bac si chua xac minh - om-H §4"),
    ("dr.", "L01 L02", "ten bac si chua xac minh - om-H §4"),
    ("bac si harvey", "L01 L02", "ten bac si chua xac minh - om-H §4"),
    ("bac si nicolas", "L01 L02", "ten bac si chua xac minh - om-H §4"),
]
CO_VANG = [
    ("tot nhat", "L04", "khang dinh tuyet doi -> 'phu hop voi tinh trang cua ban'"),
    ("duy nhat", "L04", "khang dinh tuyet doi"),
    ("chi co", "L04", "khang dinh tuyet doi -> 'vi sao bac si thuong chon...'"),
    ("so 1", "L04", "khang dinh tuyet doi"),
    ("khong dau", "L03", "chi duoc la loi KH ke, bo khoi text overlay"),
    ("khong sung", "L03", "chi duoc la loi KH ke, bo khoi text overlay"),
    ("chan dong", "L04", "tu giat gan -> bo chu nay"),
    ("surgiform", "L02", "ten thuong mai vat lieu -> noi muc loai vat lieu"),
    ("cach day 50 nam", "L04", "ha thap phuong phap khac"),
    ("mo phong", "L07", "phai chen 'Hinh mo phong, chi de tham khao'"),
]

NAC_RE = re.compile(r"\bA([1-5])\b")


# ---------------------------------------------------------------- tien ich

def bo_dau(s: str) -> str:
    """Bo dau tieng Viet, ve chu thuong, gom khoang trang."""
    if not s:
        return ""
    s = s.replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def suy_ra(v) -> str:
    return f"{v} (suy ra)"


# ---------------------------------------------------------------- doc file

def doc_csv(data: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(data)))


def doc_xlsx(path: Path, ten_sheet: str | None) -> list[list[str]]:
    try:
        import openpyxl
    except ImportError:
        sys.exit("Thieu openpyxl. Chay: pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[ten_sheet] if ten_sheet else wb[wb.sheetnames[0]]
    return [["" if c is None else str(c) for c in row]
            for row in ws.iter_rows(values_only=True)]


def tai_sheet(sheet_id: str, gid: str) -> list[list[str]]:
    url = (f"https://docs.google.com/spreadsheets/d/{sheet_id}"
           f"/export?format=csv&gid={gid}")
    with urllib.request.urlopen(url, timeout=60) as r:  # noqa: S310
        return doc_csv(r.read().decode("utf-8"))


# ---------------------------------------------------------------- map cot

COT_CAN = {
    "chan dung": "chan_dung",
    "nac 5a": "nac",
    "insight": "insight",
    "hook 3 giay": "hook",
    "cta cuoi video": "cta",
    "loai hook": "loai_hook",
    "loai du lieu": "format",
    "duyet": "duyet",
    "deadline": "deadline",
    "ten qc": "ten_qc",
    "link": "link",
    "camp dich": "camp",
    "khu vuc": "khu_vuc",
    "hieu qua": "hieu_qua",
    "ngay order": "ngay_order",
}


def tim_dong_tieu_de(rows: list[list[str]]) -> int:
    """Dong tieu de la dong co nhieu ten cot khop nhat."""
    diem_tot, idx = 0, -1
    for i, row in enumerate(rows[:40]):
        hop = " | ".join(bo_dau(c) for c in row)
        diem = sum(1 for k in COT_CAN if k in hop)
        if diem > diem_tot:
            diem_tot, idx = diem, i
    if diem_tot < 3:
        sys.exit("Khong tim duoc dong tieu de. Kiem lai file order.")
    return idx


def map_cot(header: list[str]) -> dict[str, int]:
    """Ten cot -> chi so. Giu lan khop DAU TIEN."""
    out: dict[str, int] = {}
    for i, raw in enumerate(header):
        h = bo_dau(raw)
        if not h:
            continue
        for khoa, ten in COT_CAN.items():
            if khoa in h and ten not in out:
                out[ten] = i
    return out


# ---------------------------------------------------------------- chuan hoa

def chuan_chan_dung(v: str) -> tuple[str, str]:
    b = bo_dau(v)
    if not b:
        return "", ""
    for khoa, (ma, ten) in CHAN_DUNG.items():
        if khoa in b:
            dv = v.split("-")[0].strip() if "-" in v else ""
            return ma, f"{ma} · {dv} - {ten}" if dv else f"{ma} · {ten}"
    return "", v.strip()


def chuan_nac(v: str) -> str:
    m = NAC_RE.search(bo_dau(v).upper().replace("A", "A"))
    if m:
        return f"A{m.group(1)}"
    m = re.search(r"a\s*([1-5])", bo_dau(v))
    return f"A{m.group(1)}" if m else ""


def chuan_camp(v: str) -> str:
    b = bo_dau(v)
    if not b:
        return ""
    if "view" in b:
        return "VIEW"
    if "per" in b or "perf" in b:
        return "PERF"
    return ""


def chuan_format(v: str) -> str:
    b = bo_dau(v)
    m = re.search(r"\bl(\d{2})\b", b)
    if m:
        return f"L{m.group(1)}"
    for khoa, ma in TEN_FORMAT.items():
        if khoa in b:
            return ma
    return ""


# ---------------------------------------------------------------- cham hook (B4 so bo)

def cham_hook(hook: str) -> tuple[str, str, str]:
    """-> (cham, thanh_phan_co, kieu_kem). Cham so bo de agent soi lai."""
    if not hook.strip():
        return "-", "", "thieu hook"
    b = bo_dau(hook)

    kem: list[str] = []
    for ma, ten, mau in KIEU_KEM:
        for m in mau:
            trung = (re.search(m, b) if m.startswith(("^", r"\b")) or "\\" in m
                     else (m in b))
            if trung:
                kem.append(f"{ma} {ten}")
                break

    tp: list[str] = []
    if any(t in b for t in TU_TINH_TRANG):
        tp.append("1")
    if any(t in b for t in TU_KY_THUAT):
        tp.append("2")
    if any(t in b for t in TU_BANG_CHUNG):
        tp.append("3")

    if len(b.split()) <= 4 and "1" not in tp:
        kem.append("H4 qua ngan, khong goi tinh trang")
    if not tp and not kem:
        kem.append("H3 hook mo ho")

    kem = list(dict.fromkeys(kem))
    # om-C §5: thanh phan (1) goi dung tinh trang la CONG.
    # Co (1) du mot minh -> quay duoc. Trung kieu KEM nhung van co (1) -> chi la SUA.
    if kem:
        cham = "! SUA ROI QUAY" if "1" in tp else "X DOI HOOK"
    elif "1" in tp or len(tp) >= 2:
        cham = "V QUAY DUOC"
    elif tp:
        cham = "! SUA ROI QUAY"
    else:
        cham = "X DOI HOOK"
    return cham, "".join(f"({t})" for t in tp), " · ".join(kem)


# ---------------------------------------------------------------- co phap che (B6 so bo)

def soat_phap_che(*truong: str) -> tuple[str, str, str]:
    b = bo_dau(" | ".join(t for t in truong if t))
    do = [(k, l, g) for k, l, g in CO_DO if k in b]
    vang = [(k, l, g) for k, l, g in CO_VANG if k in b]
    if do:
        return ("DO", " · ".join(l for _, l, _ in do),
                " · ".join(f"'{k}': {g}" for k, _, g in do))
    if vang:
        return ("VANG", " · ".join(l for _, l, _ in vang),
                " · ".join(f"'{k}': {g}" for k, _, g in vang))
    return "XANH", "", ""


# ---------------------------------------------------------------- xu ly mot dong

def xu_ly(r: dict[str, str], stt: int, nhan_khoi: str, dong_goc: int) -> dict[str, str]:
    cd_ma, cd_ten = chuan_chan_dung(r.get("chan_dung", ""))
    nac = chuan_nac(r.get("nac", ""))
    camp = chuan_camp(r.get("camp", ""))
    hook = (r.get("hook") or "").strip()
    cta = (r.get("cta") or "").strip()

    # suy ra
    nhiet = NHIET.get(nac, "")
    camp_sr = camp or ({"A1": "VIEW", "A5": "VIEW"}.get(nac, "PERF") if nac else "")
    o = O_CAMP.get((nac, camp_sr), "")

    fmt = chuan_format(r.get("format", ""))
    canh_bao: list[str] = []
    if o == "CHAN":
        canh_bao.append("CHAN A1+PERF: doi sang V1 hoac day nac len A2 (om-F §1)")
    if fmt in FORMAT_CHUA_DINH_NGHIA:
        canh_bao.append(f"{fmt} chua dinh nghia (om-D §1) - de xuat format thay the")
    # om-F §1: format cua o la UU TIEN, khong phai danh sach dong -> chi ghi chu.
    ghi_chu_fmt = ""
    if (fmt and fmt not in FORMAT_CHUA_DINH_NGHIA
            and FORMAT_THEO_O.get(o) and fmt not in FORMAT_THEO_O[o]):
        ghi_chu_fmt = (f"{fmt} ngoai format uu tien cua o {o} "
                       f"(uu tien: {' '.join(FORMAT_THEO_O[o])}) - van chay duoc")

    ten_fmt, thoi_luong, boi_canh, chi_so, phut, can_kh = FORMAT.get(
        fmt, ("", "", "", "", "", False))
    dao_cu = DAO_CU.get(fmt, "")
    nhan_vat = NHAN_VAT.get(fmt, "")
    tep_vao, tep_tru = TEP.get(o, ("", ""))
    nguong_tot, ket_qua = NGUONG.get(o, ("", ""))

    # Chon module chan dung theo dich vu: MAT -> om-J, con lai -> om-C (mui)
    la_mat = "mat" in bo_dau(r.get("chan_dung", ""))
    if la_mat:
        noi_so, td_thang = CHAN_DUNG_MAT.get(cd_ma, ("", ""))
        cta_mac_dinh = CTA_NAC_MAT.get(nac, "")
        if cd_ma == "CD5":
            canh_bao.append("CD5 MAT DANG LECH khoi quy uoc (mui=phong thuy) - om-J §1")
        if cd_ma in ("CD5", "CD6"):
            canh_bao.append("Cau ve BENH NEN / SUP MI phai do BS duyet - om-J §2")
    else:
        noi_so, td_thang = CHAN_DUNG_TRI_THUC.get(cd_ma, ("", ""))
        cta_mac_dinh = CTA_NAC.get(nac, "")

    # cong A
    cong_a = {
        "A1": "V" if cd_ma else "X",
        "A2": "V" if nac else "X",
        "A3": "V" if hook else "X",
        "A4": "V" if cta else "X",
        "A5": "V" if camp else "~",
        "A6": "V" if (r.get("khu_vuc") or "").strip() else "~",
        "A7": "CONSENT?" if can_kh else "-",
        "A8": "V" if (r.get("duyet") or "").strip() else "X",
    }
    chan = [k for k in ("A1", "A2") if cong_a[k] == "X"]

    cham, tp, kem = cham_hook(hook)
    co, diem_l, ghi_chu_pc = soat_phap_che(hook, cta, r.get("ten_qc", ""))

    dv = (r.get("chan_dung", "").split("-")[0].strip() or "?")
    ma_dv = bo_dau(dv)[:3].upper() or "XXX"
    ma_kb = f"KB-{ma_dv}{cd_ma[-1] if cd_ma else '?'}-{nac or 'A?'}-{stt:03d}"

    return {
        # (1) dinh danh
        "MA KICH BAN": ma_kb,
        "NGUON ORDER": f"{nhan_khoi} · dong {dong_goc}",
        "NGAY ORDER": r.get("ngay_order", ""),
        "DEADLINE QUAY": r.get("deadline", "") or "[CHO XAC NHAN]",
        "DEADLINE LEN SONG": (f"{r['deadline']} + 2 ngay dung (suy ra)"
                              if r.get("deadline") else "[CHO XAC NHAN]"),
        "TEN QC": r.get("ten_qc", ""),
        "LINK (tai dung?)": r.get("link", ""),
        # (2) chien luoc
        "DICH VU": dv,
        "CHAN DUNG": cd_ten,
        "NAC 5A": nac,
        "NHIET": suy_ra(nhiet) if nhiet else "",
        "O CAMP": suy_ra(o) if o and o != "CHAN" else ("CHAN" if o == "CHAN" else ""),
        "TEP VAO": suy_ra(tep_vao) if tep_vao else "",
        "TEP LOAI TRU": suy_ra(tep_tru) if tep_tru else "",
        "KHU VUC": r.get("khu_vuc", "") or suy_ra("CA 2"),
        "INSIGHT": r.get("insight", ""),
        "NOI SO CHI PHOI": suy_ra(noi_so) if noi_so else "",
        "THONG DIEP THANG": suy_ra(td_thang) if td_thang else "",
        # (3) kich ban
        "FORMAT": f"{fmt} · {ten_fmt}" if ten_fmt else fmt,
        "THOI LUONG": suy_ra(thoi_luong) if thoi_luong else "",
        "LOAI HOOK": r.get("loai_hook", ""),
        "CHAM HOOK": cham,
        "THANH PHAN CO": tp,
        "KIEU KEM TRUNG": kem,
        "HOOK 0-3s": hook,
        "HOOK GOC -> DE XUAT -> VI SAO": ("<< B4 — agent dien neu doi hook >>"
                                          if cham.startswith(("X", "!")) else ""),
        "CTA": cta or (suy_ra(cta_mac_dinh) if cta_mac_dinh else ""),
        "LOI THOAI THEO MOC GIAY": "<< B5 — agent viet >>",
        "TEXT ON SCREEN": "<< B5 — agent viet >>",
        "SHOT LIST + B-ROLL": "<< B5 — agent viet >>",
        "NHAC & TOC DO CAT": "<< B5 — agent viet >>",
        "CAPTION DANG": "<< B5 — agent viet >>",
        # (4) san xuat
        "BOI CANH": suy_ra(boi_canh) if boi_canh else "",
        "NHAN VAT": suy_ra(nhan_vat) if nhan_vat else "",
        "DAO CU": suy_ra(dao_cu) if dao_cu else "",
        "CAN KH THAT / CONSENT": "Co - ⛔ KIEM CONSENT" if can_kh else "Khong",
        "UOC PHUT QUAY": suy_ra(phut) if phut else "",
        "DONG BATCH": suy_ra("BATCH NHANH" if fmt in ("L01", "L02", "L11")
                             else "BATCH LON") if fmt else "",
        # (5) cong chan
        "CONG A": " ".join(f"{k}:{v}" for k, v in cong_a.items()),
        "CHAN": " + ".join(chan) if chan else "",
        "CO PHAP CHE": co,
        "DIEM L": diem_l,
        "CAU PHAI SUA -> THAY THE": ghi_chu_pc,
        "AI DUYET": r.get("duyet", "") or "CHUA CO NGUOI DUYET",
        # (6) do luong
        "CHI SO DO CHINH": suy_ra(chi_so) if chi_so else "",
        "KET QUA DO": suy_ra(ket_qua) if ket_qua else "",
        "NGUONG TOT": suy_ra(nguong_tot) if nguong_tot else "",
        "NGUONG KEM": suy_ra(">1,3x nguong tot") if nguong_tot else "",
        "HIEU QUA (ad cu)": r.get("hieu_qua", ""),
        "CANH BAO": " | ".join(canh_bao),
        "GHI CHU FORMAT": ghi_chu_fmt,
    }


# ---------------------------------------------------------------- main

def main() -> None:
    # Console Windows mac dinh cp1252, khong in duoc tieng Viet.
    for s in (sys.stdout, sys.stderr):
        if hasattr(s, "reconfigure"):
            s.reconfigure(encoding="utf-8", errors="replace")

    p = argparse.ArgumentParser(description="Doc & chuan hoa file order noi dung")
    p.add_argument("file", nargs="?", help="file .csv hoac .xlsx")
    p.add_argument("--sheet", help="ten sheet trong file xlsx")
    p.add_argument("--sheet-id", help="Google Spreadsheet ID")
    p.add_argument("--gid", help="gid cua tab")
    p.add_argument("--out", default="out/phieu-kich-ban.csv")
    a = p.parse_args()

    if a.sheet_id and a.gid:
        rows, nguon = tai_sheet(a.sheet_id, a.gid), f"gsheet:{a.gid}"
    elif a.file:
        path = Path(a.file)
        if not path.exists():
            sys.exit(f"Khong thay file: {path}")
        rows = (doc_xlsx(path, a.sheet) if path.suffix.lower() in (".xlsx", ".xlsm")
                else doc_csv(path.read_text(encoding="utf-8-sig")))
        nguon = path.name
    else:
        p.print_help()
        sys.exit(1)

    hi = tim_dong_tieu_de(rows)
    cot = map_cot(rows[hi])
    thieu = {"chan_dung", "nac", "hook"} - set(cot)
    if thieu:
        print(f"CANH BAO: khong tim thay cot {thieu} — ket qua se khuyet.",
              file=sys.stderr)

    ket: list[dict[str, str]] = []
    nhan_khoi, bo_qua, stt = nguon, 0, 0

    for i, row in enumerate(rows[hi + 1:], start=hi + 2):
        if not any(c.strip() for c in row):
            continue
        r = {ten: (row[idx].strip() if idx < len(row) else "")
             for ten, idx in cot.items()}

        # dong tieu de khoi: chi 1-2 o co chu, khong co hook
        dac = [c for c in row if c.strip()]
        if len(dac) <= 2 and not r.get("hook"):
            nhan_khoi = f"{nguon} · {dac[0].strip()}" if dac else nhan_khoi
            continue
        if not any(r.get(k) for k in ("chan_dung", "nac", "hook", "insight")):
            bo_qua += 1
            continue

        stt += 1
        ket.append(xu_ly(r, stt, nhan_khoi, i))

    if not ket:
        sys.exit("Khong doc duoc dong nao co noi dung.")

    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(ket[0].keys()))
        w.writeheader()
        w.writerows(ket)

    # ---- bao cao
    chan = [r for r in ket if r["CHAN"]]
    do = [r for r in ket if r["CO PHAP CHE"] == "DO"]
    vang = [r for r in ket if r["CO PHAP CHE"] == "VANG"]
    doi = [r for r in ket if r["CHAM HOOK"].startswith("X")]
    sua = [r for r in ket if r["CHAM HOOK"].startswith("!")]
    cb = [r for r in ket if r["CANH BAO"]]
    consent = [r for r in ket if r["CAN KH THAT / CONSENT"].startswith("Co")]

    print(f"\n  DOC DUOC : {len(ket)} dong   ·   BO QUA: {bo_qua} dong")
    print(f"  GHI RA   : {out}\n")
    print(f"  CONG A — CHAN (thieu chan dung / nac 5A) : {len(chan)}")
    for r in chan[:12]:
        print(f"      {r['MA KICH BAN']}  [{r['CHAN']}]  {r['HOOK 0-3s'][:52]}")
    print(f"\n  B4 SOI HOOK — X DOI HOOK: {len(doi)}   ! SUA ROI QUAY: {len(sua)}"
          f"   V QUAY DUOC: {len(ket) - len(doi) - len(sua)}")
    for r in doi[:12]:
        print(f"      {r['MA KICH BAN']}  {r['KIEU KEM TRUNG'][:40]:40} "
              f"{r['HOOK 0-3s'][:44]}")
    print(f"\n  B6 PHAP CHE — DO: {len(do)}   VANG: {len(vang)}   "
          f"XANH: {len(ket) - len(do) - len(vang)}")
    for r in (do + vang)[:12]:
        print(f"      [{r['CO PHAP CHE']:5}] {r['MA KICH BAN']}  "
              f"{r['CAU PHAI SUA -> THAY THE'][:66]}")
    if cb:
        print(f"\n  CANH BAO — CHAN : {len(cb)}")
        for r in cb[:12]:
            print(f"      {r['MA KICH BAN']}  {r['CANH BAO'][:76]}")
    gc = [r for r in ket if r["GHI CHU FORMAT"]]
    if gc:
        print(f"\n  GHI CHU FORMAT (khong chan, chi de biet) : {len(gc)} dong")
    if consent:
        print(f"\n  CAN CONSENT (format co KH that) : {len(consent)}")
        for r in consent[:12]:
            print(f"      {r['MA KICH BAN']}  {r['FORMAT']}")
    print("\n  >> Buoc tiep: dua file nay cho agent chay B4-B7 "
          "(soi lai hook, viet loi thoai, call sheet).\n")


if __name__ == "__main__":
    main()
